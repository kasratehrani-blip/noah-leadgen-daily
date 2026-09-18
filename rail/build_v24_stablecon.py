#!/usr/bin/env python3
"""Stablecon attendees as a v24-style lead-gen workbook, one tab per Noah vertical.
Usage: build_v24_stablecon.py <ranked.json> <hs_desc.json> <out.xlsx>"""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
src, hsd, out = sys.argv[1:4]
J = json.load(open(src)); HS = json.load(open(hsd))
OWN = {'78072891':'Loek De Jager','78072893':'Charlie Bunn','78072890':'Dylan Rayapah','78072935':'Thijn Lamers','35955904':'Faris Riaz','32020220':'Jameson Lobb','30544029':'Casimir Hohenlohe','78307389':'Nikrad Nassiri','43643762':'Shah Ramezani','78072892':'owner 78072892','36176637':'owner 36176637'}
DATE = '18 Sep 2026'
FIXED = ("Noah is the stablecoin infrastructure powering modern payments, with the ability to issue named USD, EUR and GBP virtual accounts in 160+ markets, and local payout in 70+ markets. "
         "We help move money worldwide faster and more cheaply than correspondent banking - powering the likes of Deel, Toptal, WorldRemit and Ledger.")
INTRO = ("I hope you are doing well. My name is Thijn Lamers, co-founder of Noah and before that part of the founding team at Adyen. "
         "It would be great if we could schedule a call, even 15 minutes, to explore how we could work together and help each other.")
OUTRO = "Looking forward to hearing from you.\n\nMany thanks,\n\nThijn"

def nm(s):
    return ' '.join(w.capitalize() if len(w) > 1 else w for w in s.split()) if (s.islower() or s.isupper()) else s
def short(c):
    c = re.sub(r'\s*\(.*?\)\s*', ' ', c).strip()
    c = re.sub(r',?\s+(Inc\.?|LLC|Ltd\.?|Limited|Corp\.?|S\.A\.|SA|AG|GmbH|Pte\.? Ltd\.?|LLP)$', '', c, flags=re.I).strip()
    return c
def is_cu(c): return bool(re.search(r'credit union|\bFCU\b|\bCU\b', c, re.I))
def size_band(e):
    try: e = int(e)
    except: return ''
    return '1-10' if e <= 10 else '11-50' if e <= 50 else '51-200' if e <= 200 else '201-1000' if e <= 1000 else '1000+'
def mmc(seg, emp, cu=False):
    try: e = int(emp or 0)
    except: e = 0
    if seg == 'RAILS': return '$3-10k entry' if e and e <= 50 else '$10-30k tier' if e and e <= 250 else '$30k+ / structured partnership' if e else '$5-20k tier (size unverified)'
    if seg == 'BANKS': return '$3-10k entry (member remittance)' if cu else '$5-20k tier, volume-led'
    if seg == 'CRYPTO': return 'Partnership / referral; $5-15k if embedded'
    if seg == 'PAYMENTS': return 'Unverified: classify first'
    return 'n/a (not a rails buyer)'
SEG = {5:'RAILS', 4:'BANKS', 3:'CRYPTO', 2:'VENDORS', 1:'MEDIA'}
def seg_of(r):
    if r['clbl'].startswith('Payments-shaped'): return 'PAYMENTS'
    if r['clbl'] == 'Other': return 'OTHER'
    return SEG[r['cf']]
def hs_status(r):
    h = r['hubspot']
    if h == 'New': return 'No - net-new'
    m = re.search(r'\(([^)]*)\)', h); inner = m.group(1) if m else ''
    owner = next((p.strip() for p in inner.split(',') if re.match(r'^[A-Z][a-z]+ [A-Z]', p.strip())), '')
    stage = 'deal on record' if 'deal' in inner else ('opportunity' if 'opportunity' in inner else 'no deal on record')
    if h.startswith('Contact'): return f"Yes - contact owned by {owner or 'someone at Noah'} - coordinate first"
    return f"Yes - company known ({stage}{', ' + owner if owner else ''}) - person is new"
def route(r, seg):
    if seg in ('VENDORS',): return 'PARTNER (channel, not a buyer)'
    if seg == 'MEDIA': return 'RELATIONSHIP / PR (no product pitch)'
    if seg == 'OTHER': return 'WATCH (classify first)'
    if r['hsk'] == 'contact': return 'SELL via owner (coordinate first)'
    if r['rf'] == 3: return 'SELL - write first'
    if r['rf'] == 2: return 'SELL - write second'
    return 'SELL - later (find the payments owner)'

def copy_for(r, seg):
    C = short(r['company']); cu = is_cu(r['company'])
    who = 'members' if cu else 'customers'
    if seg in ('RAILS', 'PAYMENTS'):
        return dict(
            fit=f"{C} moves money across borders and needs fiat legs. Noah gives them named USD, EUR and GBP virtual accounts in 160+ markets to collect, stablecoin settlement in the middle, and local payout in 70+ markets - so they stop stitching bank partners per corridor." + (" Segment inferred from the company name; confirm what they do before drafting." if seg == 'PAYMENTS' else ''),
            flow="USD/EUR/GBP in via named VAs (ACH, Fedwire, SWIFT, SEPA, FPS) -> stablecoin settlement -> local-rail payout in 70+ markets || VAs: USD + EUR + GBP. Payouts: business and individual.",
            pain="Every new corridor is a new bank partner or licence; correspondent banking cost and float in transit; named-account coverage outside the home market; reconciliation across multiple providers.",
            incumbent="Own banking partners plus one of Bridge, BVNK or Circle for the stablecoin leg; SWIFT correspondents for the rest. Noah replaces the fiat legs, not their product. Pricing unverified.",
            model="Partner-licensed or own licence; Standard KYB for business flows; Reliance viable where they hold licences.",
            hook=f"named USD, EUR and GBP accounts and local payout behind {C}'s cross-border flows",
            subject="Noah for your fiat legs")
    if seg == 'BANKS':
        return dict(
            fit=f"{C}'s {who} send money abroad through correspondent wires. Noah gives a bank or credit union local payout in 70+ markets and named USD, EUR and GBP accounts without building corridors - a complement to the core, not a core replacement.",
            flow=f"{who.capitalize()} USD out -> stablecoin settlement -> local payout in 70+ markets || VAs: USD (+EUR/GBP for business clients). Payouts: individuals and SME suppliers.",
            pain=f"International wires via correspondents are slow, expensive and opaque; {who} leave for Wise or Remitly; FX and correspondent fees erode margin; SME clients paying overseas suppliers.",
            incumbent="Correspondent banks over SWIFT, sometimes a white-label remittance partner. Typical 0.5-1% plus fixed fee vs Noah bps.",
            model="Regulated FI: Reliance-shaped viable; bank runs its own KYC, Noah is the rail.",
            hook=f"local payout in 70+ markets when {C}'s {who} send money abroad",
            subject=f"Noah for your {who} abroad")
    if seg == 'CRYPTO':
        return dict(
            fit=f"{C} sells custody, wallets or chain infrastructure; their clients still need fiat on and off legs. Noah supplies named USD/EUR/GBP accounts plus local payout as the fiat layer under their stack - partner as much as customer.",
            flow="Client fiat in via named VA -> stablecoin on their rails -> fiat out via local payout in 70+ markets || VAs: USD + EUR + GBP.",
            pain="Fiat legs are the gap: bank partners per market, ramps are consumer-grade and patchy, clients ask them for a business-grade fiat rail.",
            incumbent="Consumer ramps (MoonPay, Transak) or none. Noah is the B2B fiat layer they can embed or refer.",
            model="Partner-rail; Token Share KYB where they onboard the client.",
            hook=f"the fiat legs under {C}'s stablecoin stack",
            subject="Noah under your stack")
    if seg == 'VENDORS':
        return dict(
            fit=f"{C} sells compliance, identity or analytics to the same fintechs Noah sells to. Channel and referral partner, not a rails buyer.",
            flow="n/a - channel partner.",
            pain="Their clients ask who runs the fiat rails; a referral answer helps them close.",
            incumbent="n/a.", model="n/a.",
            hook=f"the fiat legs {C}'s clients keep asking about",
            subject="Noah alongside " + C)
    if seg == 'MEDIA':
        return dict(fit=f"{C} is media, events, research or advisory. Relationship and PR value; no product pitch.", flow="n/a.", pain="n/a.", incumbent="n/a.", model="n/a.", hook="", subject="")
    return dict(fit="Not classified: check what the company does before any outreach.", flow="", pain="", incumbent="", model="", hook="", subject="")

def angle(r, seg):
    C = short(r['company'])
    if seg in ('MEDIA', 'OTHER'): return 'No cold pitch. Thijn personalises only if he knows them.'
    if seg == 'VENDORS': return f"Partner note: we both sell to stablecoin fintechs; offer a referral loop, not a rails pitch."
    if r['rlbl'] == 'Decision maker': return f"Founder/CEO decides. Lead with the corridor they cannot cover today and the {DATE.split()[1]} Stablecon meet as the hook; one call, 15 minutes."
    if r['rlbl'] == 'Buyer of rails': return f"They own payments or partnerships at {C}. Lead with coverage map: which of their markets Noah already pays out to, and the named-account gap."
    if r['rlbl'].startswith('Senior'): return f"Senior but not payments-specific. Ask for the intro to whoever owns payments at {C}; keep the note short."
    return f"Junior or unclear role at {C}. Use them to find the payments owner; do not pitch."

def draft(r, seg, c):
    if not c['hook']: return '', '(no draft - relationship or classify first)'
    first = nm(r['name']).split()[0]; C = short(r['company'])
    body = (f"Hi {first},\n\n{INTRO}\n\nAbout Noah\n{FIXED}\n\nWe are the missing piece for {c['hook']}, and I'd love the opportunity to explore it with {C}.\n\n{OUTRO}")
    return 'Re: ' + c['subject'], body

HDR = ['HubSpot (status - owner)','Route','Rank','Score','Company','Domain','Country / Region','Size (emp)','What they do','Why this is a Noah fit','Payment flow (currency + payout mode)','Likely pain','Expected MMC (est.)','Buyer & key contact','Title','Role fit','Email','LinkedIn','Outreach angle','Incumbent supplier vs Noah','Likely compliance model','Draft subject','Draft body','Source','Date checked']
W = [30,22,6,6,24,20,18,9,48,52,44,40,20,22,30,16,30,34,44,40,30,26,70,26,12]
WRAP = {8,9,10,11,13,18,19,20,22}
def row_for(r):
    seg = seg_of(r); h = HS.get(r['domain'], {}); c = copy_for(r, seg)
    region = ', '.join(x for x in (h.get('city',''), h.get('country','')) if x) or 'verify'
    what = h.get('desc') or f"Not in HubSpot. Segment: {r['clbl']}. Verify on the web or enrich in Apollo before drafting."
    subj, body = draft(r, seg, c)
    return [hs_status(r), route(r, seg), r.get('rank',''), r.get('score',''), r['company'], r['domain'], region, size_band(h.get('emp','')), what, c['fit'], c['flow'], c['pain'], mmc(seg, h.get('emp',''), is_cu(r['company'])), nm(r['name']), r['title'], r['rlbl'], r['email'], r['linkedin'], angle(r, seg), c['incumbent'], c['model'], subj, body, "Stablecon badge wall (Faris, 9 Sep 2026); Apollo people match 16 Sep 2026; HubSpot check 16 Sep 2026", DATE], seg, r

F = lambda **k: Font(name='Arial', **{'size': 10, **k})
HFILL = PatternFill('solid', fgColor='1F3864'); GREEN = PatternFill('solid', fgColor='E2F0D9'); AMBER = PatternFill('solid', fgColor='FFF2CC'); BLUE = PatternFill('solid', fgColor='DDEBF7'); GREY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='D9D9D9')
wb = Workbook(); wb.remove(wb.active)
def sheet(name, note, rows, hdr=HDR, widths=W, wrap=WRAP, colour=True, freeze='F3'):
    ws = wb.create_sheet(name)
    ws['A1'] = note; ws['A1'].font = F(italic=True); ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(hdr), 12)); ws.row_dimensions[1].height = 42
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=2, column=i, value=h); c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 30
    for j, (vals, seg, r) in enumerate(rows, 3):
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=j, column=i, value=v); c.font = F(); c.border = Border(bottom=thin); c.alignment = Alignment(wrap_text=(i - 1) in wrap, vertical='top')
        if colour and r is not None:
            fill = AMBER if r.get('hsk') == 'contact' else GREEN if r.get('hsk') == 'new' else None
            if fill: ws.cell(row=j, column=1).fill = fill
            if str(vals[1]).startswith('SELL - write first'): ws.cell(row=j, column=2).fill = BLUE
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze; ws.auto_filter.ref = f"A2:{get_column_letter(len(hdr))}{len(rows) + 2}"
    return ws

ranked = J['ranked']; inv = J['investors']; cust = J['customers']
rows = [row_for(r) for r in ranked]
by = {}
for x in rows: by.setdefault(x[1], []).append(x)
TABS = [('RAILS', 'Cross-border payments, stablecoin rails, remittance, payouts, neobanks, ramps and exchanges. Noah core ICP: they need named USD/EUR/GBP VAs and local payout in 70+ markets. Row order = rank.'),
        ('BANKS', 'Banks, credit unions, card issuers and payment networks. Angle: local payout in 70+ markets for their customers or members, a complement to the core. Row order = rank.'),
        ('CRYPTO', 'Crypto infrastructure, custody, wallets and chains. Their clients need fiat legs; Noah is the fiat layer under the stack. Partner as much as customer. Row order = rank.'),
        ('PAYMENTS', 'Looks payments-shaped from the name but not classified. Verify what they do before drafting.'),
        ('VENDORS', 'Compliance, identity, analytics and security vendors. Channel partners, not rails buyers. No product pitch.'),
        ('MEDIA', 'Media, events, research and advisory. Relationship and PR only.'),
        ('OTHER', 'Not classified. Check before any outreach.')]
def kpi(seg): xs = by.get(seg, []); return len(xs), sum(1 for x in xs if x[2]['hsk'] == 'new'), sum(1 for x in xs if x[0][1].startswith('SELL - write first'))
# START HERE
ws = wb.create_sheet('START HERE')
lines = [('Noah lead gen - Stablecon attendees by vertical - ' + DATE, True),
         ('What this is', True), ('Every Stablecon attendee (1,041 badges from Faris, 9 Sep 2026) run through Apollo for current company and a verified work email, checked against HubSpot, scored against Noah ICP and split by vertical. One row per person, v24 column standard.', False),
         ('', False), ('Tabs', True)]
for seg, note in TABS:
    n, new, first = kpi(seg); lines.append((f"{seg}: {n} people, {new} net-new to Noah, {first} write first. {note}", False))
lines += [(f"INVESTORS: {len(inv)} people. Shah decides; no product pitch.", False), (f"CUSTOMERS: {len(cust)} people at companies already live with Noah. Relationship, not outreach.", False), ('ALL RANKED: every person in one list, rank order.', False), ('', False),
          ('How to read a row', True),
          ('HubSpot column: green = net-new (nobody at Noah has the person or company); amber = the contact is owned by someone at Noah, coordinate first; white = the company is known but this person is new.', False),
          ('Route: SELL write first = decision maker or buyer of rails at a net-new or company-known account; write second = senior but not payments-specific; later = junior, use them to find the owner. SELL via owner = amber. PARTNER, RELATIONSHIP, WATCH as labelled. Blue = write first.', False),
          ('Score = company fit x3 (5 rails, 4 bank/network, 3 crypto infra, 2 vendor, 1 media) + role fit (3 decision maker or rails buyer, 2 senior, 1 other) + HubSpot (+1 new, 0 company known, -1 owned). Max 19.', False),
          ('What they do: HubSpot description where the company is in the CRM; otherwise marked for verification (Apollo enrich or web).', False),
          ('Draft subject and body: Thijn standard (intro, About Noah, fixed paragraph, hook line, outro), subject prefixed Re:. Hook is per vertical; personalise with the What they do column before loading.', False),
          ('Expected MMC: estimate by vertical and size, unverified.', False), ('', False),
          ('Rules', True), ('Nothing sends from this book. HubSpot is read-only. Confirm title on the web before drafting; Apollo titles were not individually checked. Amber rows go through the owner.', False)]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = F(bold=b, size=12 if i == 1 else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 140
for seg, note in TABS:
    if by.get(seg): sheet(seg, note, by[seg])
IH = ['HubSpot (status - owner)','Route','Name','Title','Company','Domain','Country / Region','What they do','Email','LinkedIn','Angle','Source','Date checked']
IW = [30,26,24,30,26,20,18,48,30,34,44,26,12]
def inv_row(o):
    h = HS.get(o['domain'], {}); region = ', '.join(x for x in (h.get('city',''), h.get('country','')) if x) or 'verify'
    return [hs_status(o), 'INVESTOR (Shah decides)', nm(o['name']), o['title'], o['company'], o['domain'], region, h.get('desc') or 'Not in HubSpot; verify.', o['email'], o['linkedin'], 'Intro path only. Thijn or Shah personalises if the relationship is real. Not a product pitch.', 'Stablecon badge wall; Apollo 16 Sep 2026; HubSpot 16 Sep 2026', DATE], 'INV', o
sheet('INVESTORS', 'Investors and funds on the badge wall with a verified email. Shah decides who to approach. No product pitch.', [inv_row(o) for o in inv], IH, IW, {7,10}, freeze='D3')
def cust_row(o):
    h = HS.get(o['domain'], {}); region = ', '.join(x for x in (h.get('city',''), h.get('country','')) if x) or 'verify'
    return ['Yes - CUSTOMER - ' + (o['hubspot'] or ''), 'RELATIONSHIP (already live)', nm(o['name']), o['title'], o['company'], o['domain'], region, h.get('desc') or '', o['email'], o['linkedin'], 'Say hello, no pitch. Route through the account owner.', 'Stablecon badge wall; Apollo 16 Sep 2026; HubSpot 16 Sep 2026', DATE], 'CUST', o
sheet('CUSTOMERS', 'Attendees from companies already live with Noah. Relationship, not outreach.', [cust_row(o) for o in cust], IH, IW, {7,10}, colour=False, freeze='D3')
sheet('ALL RANKED', f'All {len(rows)} ranked people in one list, rank order. Same columns as the vertical tabs.', rows)
# NOTES
wn = wb.create_sheet('NOTES')
notes = ['Column standard follows Noah Quick-Sign v24 (Jul 2026): HubSpot status, company, domain, region, what they do, why Noah fit, payment flow, likely pain, expected MMC, buyer and contact, outreach angle, incumbent vs Noah, compliance model, draft, source, date checked.',
         'Differences from v24: one row per person (not per company) because the source is a badge wall; Role fit, Rank and Score columns added; "What they are looking for" and "Who to hunt" dropped because the contact is already identified with a verified email.',
         'Partners excluded: ATTRUS and Trace Finance (Apollo account tagged ZZZ-PARTNER).',
         'Companies not in HubSpot have no description; enrich in Apollo (1 credit per company) or check the web before drafting.',
         'Segment map lives in rail/icp_rank.py (domain lists). A company in the wrong tab = fix the map and rebuild.']
for i, t in enumerate(notes, 1):
    c = wn.cell(row=i, column=1, value=t); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical='top')
wn.column_dimensions['A'].width = 140
wb.save(out)
print({s: len(by.get(s, [])) for s, _ in TABS}, 'inv', len(inv), 'cust', len(cust), 'total', len(rows))
