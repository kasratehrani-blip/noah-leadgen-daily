#!/usr/bin/env python3
"""Owner-route workbook: Rail live clients already in HubSpot (by owner) + Stablecon contacts owned by someone at Noah. Loek first."""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
S = sys.argv[1]; out = sys.argv[2]
OWN = {'78072891':'Loek De Jager','78072893':'Charlie Bunn','78072890':'Dylan Rayapah','78072935':'Thijn Lamers','35955904':'Faris Riaz','32020220':'Jameson Lobb','30544029':'Casimir Hohenlohe','78307389':'Nikrad Nassiri','43643762':'Shah Ramezani','78072892':'owner 78072892','36176637':'owner 36176637'}
def owner(ev):
    ev = ev or ''
    for k, v in OWN.items():
        if k in ev or v.lower() in ev.lower(): return v
    return 'no owner on record'
def stage(ev):
    e = (ev or '').lower()
    if 'customer' in e: return 'Customer'
    if 'deal' in e and re.search(r'[1-9]\s*deal', e): return 'Open deal'
    if 'opportunity' in e: return 'Opportunity'
    return 'Lead'
def link(ev):
    m = re.search(r'https://app\.hubspot\.com/contacts/145864360/record/0-[123]/\d+', ev or ''); return m.group(0) if m else ''
def lastc(ev):
    m = re.search(r'last[_ ]contacted[:= ]*\s*(\d{4}-\d{2}-\d{2})', ev or '', re.I); return m.group(1) if m else ''
rows = json.load(open(S + '/results/merged.json')) + json.load(open(S + '/live2_pack.json'))
seen = set(); live = []
for r in rows:
    if (r.get('hubspot_status') or '').upper() == 'NET_NEW' or r['company'] in seen: continue
    seen.add(r['company']); ev = str(r.get('hubspot_evidence') or '')
    live.append([owner(ev), r['company'], r.get('domain',''), stage(ev), lastc(ev), r.get('ceo_name',''), r.get('ceo_title',''), r.get('ceo_email',''), r.get('email_status',''), (r.get('description') or '')[:300], (r.get('fit_hook') or ''), link(ev), ev[:400]])
live.sort(key=lambda x: (x[0] != 'Loek De Jager', x[0], x[1].lower()))
J = json.load(open('out/Stablecon_Ranked_for_Outreach_16_Sep_2026.json'))
sc = []
for o in J['ranked'] + J['investors']:
    if o.get('hsk') == 'contact':
        sc.append([owner(o['hubspot']), o['name'], o['title'], o['company'], o['domain'], o['email'], o['linkedin'], o['hubspot'], o.get('rlbl',''), o.get('score','')])
sc.sort(key=lambda x: (x[0] != 'Loek De Jager', x[0], -(x[9] or 0)))
F = lambda **k: Font(name='Arial', **{'size': 10, **k}); LINK = Font(name='Arial', size=10, color='0563C1', underline='single'); AMBER = PatternFill('solid', fgColor='FFF2CC')
wb = Workbook(); wb.remove(wb.active)
def sheet(name, note, hdr, widths, rows, wrap, links):
    ws = wb.create_sheet(name); ws['A1'] = note; ws['A1'].font = F(italic=True); ws['A1'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells('A1:H1'); ws.row_dimensions[1].height = 34
    for i, h in enumerate(hdr, 1): ws.cell(row=2, column=i, value=h)
    for j, vals in enumerate(rows, 3):
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=j, column=i, value=v); c.font = F(); c.alignment = Alignment(wrap_text=(i - 1) in wrap, vertical='top')
            if i in links and v:
                if '@' in str(v): c.hyperlink = f"mailto:{v}"
                else: c.hyperlink = str(v); c.value = 'open'
                c.font = LINK
        if vals[0] == 'Loek De Jager': ws.cell(row=j, column=1).fill = AMBER
        ws.row_dimensions[j].height = 48
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    t = Table(displayName=re.sub(r'[^A-Za-z0-9]', '', name) + 'T', ref=f"A2:{get_column_letter(len(hdr))}{len(rows) + 2}"); t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); ws.add_table(t); ws.freeze_panes = 'C3'
H1 = ['Owner at Noah','Company','Domain','HubSpot stage','Last contacted','CEO','CEO title','CEO email','Email status','What they do','Noah hook','HubSpot record','HubSpot evidence']
W1 = [18,30,20,14,13,22,26,30,14,50,40,10,60]
loek_live = [x for x in live if x[0] == 'Loek De Jager']
sheet('LOEK - live clients', f'Rail live clients that are already in HubSpot and owned by Loek. {len(loek_live)} companies. Not cold-drafted for Thijn; Loek continues the conversation. CEO name and email from Apollo (Sep 2026).', H1, W1, loek_live, {9,10,12}, {8,12})
H2 = ['Owner at Noah','Name','Title','Company','Domain','Email','LinkedIn','HubSpot','Role fit','Score']
W2 = [18,22,30,26,20,30,10,44,16,7]
loek_sc = [x for x in sc if x[0] == 'Loek De Jager']
sheet('LOEK - Stablecon', f'Stablecon attendees whose HubSpot contact is owned by Loek. {len(loek_sc)} people. Loek writes or hands over.', H2, W2, loek_sc, {7}, {6,7})
sheet('ALL - live clients by owner', f'All Rail live clients already in HubSpot, grouped by owner. {len(live)} companies. Amber = Loek.', H1, W1, live, {9,10,12}, {8,12})
sheet('ALL - Stablecon by owner', f'All Stablecon attendees already owned as a contact in HubSpot, by owner. {len(sc)} people.', H2, W2, sc, {7}, {6,7})
wb.save(out)
from collections import Counter
print('live by owner', Counter(x[0] for x in live)); print('stablecon by owner', Counter(x[0] for x in sc))
